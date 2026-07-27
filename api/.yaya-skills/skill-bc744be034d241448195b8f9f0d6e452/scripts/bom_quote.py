#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Prepare BOM search tasks and build supplier-backed quote workbooks."""

from __future__ import annotations

import argparse
import http.cookiejar
import json
import math
import re
import time
import urllib.request
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import quote_plus

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SHEET_RAW_BOM = "\u539f\u59cbbom"
HEADER_QTY_CN = "\u7528\u91cf"
HEADER_PURCHASE_QTY_CN = "\u91c7\u8d2d\u6570\u91cf"
HEADER_UNIT_CN = "\u5355\u4f4d"
WAIT_QUERY = "\u5f85\u67e5\u8be2"
SITE_LCSC = "\u7acb\u521b\u5546\u57ce"
SITE_ICKEY = "ICKEY"
SITE_DEFAULTS = [SITE_LCSC, SITE_ICKEY, "Website3", "Website4"]
LCSC_HOME_URL = "https://so.szlcsc.com/"
LCSC_API_URL = "https://so.szlcsc.com/query/product"
IN_STOCK_TOKENS = {"\u73b0\u8d27", "\u6709\u8d27", "in stock"}
OUT_STOCK_TOKENS = {"\u65e0\u8d27", "\u7f3a\u8d27", "out of stock"}
WARN_NOT_STOCK = "\u975e\u73b0\u8d27\u6216\u73b0\u8d27\u72b6\u6001\u672a\u77e5"
WARN_STOCK_SHORT = "\u5e93\u5b58\u4e0d\u8db3\u6216\u672a\u77e5"
WARN_MOQ = "\u672a\u6ee1\u8db3MOQ\u6216MOQ\u672a\u77e5"

MAIN_HEADERS_BASE = [
    "\u5e8f\u53f7",
    "\u7269\u6599\u540d\u79f0",
    "\u5236\u9020\u5546",
    "\u5236\u9020\u5546\u578b\u53f7",
    "\u89c4\u683c\u53c2\u6570",
    "\u4f4d\u53f7",
    "\u7528\u91cf",
    "\u5355\u4f4d",
]
HEADER_ORDER_QTY = "\u4e0b\u5355\u6570\u91cf\uff08\u91c7\u8d2d\u6570\u91cf\uff09"
HEADER_REMARK = "\u5907\u6ce8"
HEADER_LEAD_TIME = "\u4ea4\u671f"
SITE_SUBHEADERS = [
    "\u7f51\u5740",
    "\u547d\u4e2d\u540e\u7684bom\u7684\u5e97\u94fa\u94fe\u63a5",
    "\u67e5\u4ef7\u683c",
    "\u5e93\u5b58\u6570\u91cf",
    "MOQ\uff08\u6700\u5c0f\u8d77\u8ba2\u91cf\uff09",
    "\u67e5\u4ea4\u671f",
]
DEFAULT_QUOTE_SHEET_NAME = "\u6574\u7406\u540e\u7684\u62a5\u4ef7"

MAIN_FILL = PatternFill("solid", fgColor="D9EAF7")
SITE_FILL = PatternFill("solid", fgColor="FCE4D6")
SUB_FILL = PatternFill("solid", fgColor="E2F0D9")
THIN_BORDER = Border(
    left=Side(style="thin", color="A6A6A6"),
    right=Side(style="thin", color="A6A6A6"),
    top=Side(style="thin", color="A6A6A6"),
    bottom=Side(style="thin", color="A6A6A6"),
)


@dataclass
class BomRow:
    source_row: int
    index: int
    vision_name: str
    qty: float
    purchase_quantity: float
    total_quantity: float
    search_keyword: str
    material_name_hint: str
    unit: str


def normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "").replace("_", "")


def to_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        result = float(value)
        return None if math.isnan(result) else result
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    text = re.sub(r"[\uFFE5\u00A5$]", "", text)
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    return float(match.group(0))


def normalize_bool(value: Any) -> bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "yes", "y", "1", *IN_STOCK_TOKENS}:
        return True
    if text in {"false", "no", "n", "0", *OUT_STOCK_TOKENS}:
        return False
    return None


def split_vision_name(value: str) -> tuple[str, str]:
    text = str(value or "").strip()
    for sep in (":", "\uff1a"):
        if sep in text:
            left, right = text.split(sep, 1)
            return left.strip(), right.strip()
    return "", text


def find_source_sheet(workbook: Any, requested: str | None) -> Any:
    if requested:
        return workbook[requested]
    if SHEET_RAW_BOM in workbook.sheetnames:
        return workbook[SHEET_RAW_BOM]
    for ws in workbook.worksheets:
        headers = find_header(ws, require_purchase_quantity=False)
        if headers:
            return ws
    return workbook.worksheets[0]


def find_header(ws: Any, require_purchase_quantity: bool) -> tuple[int, dict[str, int]] | None:
    aliases = {
        "vision_name": {"visionname", "vision_name"},
        "qty": {"qty", HEADER_QTY_CN},
        "purchase_quantity": {"purchasequantity", "purchase_quantity", HEADER_PURCHASE_QTY_CN},
        "unit": {"unit", HEADER_UNIT_CN},
    }
    for row_index in range(1, min(ws.max_row, 60) + 1):
        mapping: dict[str, int] = {}
        for col_index in range(1, ws.max_column + 1):
            key = normalize_header(ws.cell(row_index, col_index).value)
            for canonical, names in aliases.items():
                if key in names and canonical not in mapping:
                    mapping[canonical] = col_index
        required = {"vision_name", "qty"}
        if require_purchase_quantity:
            required.add("purchase_quantity")
        if required.issubset(mapping):
            return row_index, mapping
    return None


def read_bom_rows(
    bom_path: Path,
    sheet_name: str | None,
    default_purchase_quantity: float | None,
) -> list[BomRow]:
    wb = load_workbook(bom_path, data_only=False, read_only=True)
    ws = find_source_sheet(wb, sheet_name)
    header = find_header(ws, require_purchase_quantity=default_purchase_quantity is None)
    if not header:
        if default_purchase_quantity is None:
            raise SystemExit(
                "Cannot find Vision_name, Qty, and Purchase_Quantity. "
                "Pass --purchase-quantity if this BOM uses one quote quantity for all rows."
            )
        header = find_header(ws, require_purchase_quantity=False)
    if not header:
        raise SystemExit("Cannot find a header row containing Vision_name and Qty.")

    header_row, columns = header
    rows: list[BomRow] = []
    for excel_row in range(header_row + 1, ws.max_row + 1):
        vision_name = str(ws.cell(excel_row, columns["vision_name"]).value or "").strip()
        if not vision_name:
            continue
        qty = to_number(ws.cell(excel_row, columns["qty"]).value)
        if qty is None:
            continue
        purchase_quantity = None
        if "purchase_quantity" in columns:
            purchase_quantity = to_number(ws.cell(excel_row, columns["purchase_quantity"]).value)
        if purchase_quantity is None:
            purchase_quantity = default_purchase_quantity
        if purchase_quantity is None:
            raise SystemExit(f"Missing Purchase_Quantity at source row {excel_row}.")

        material_hint, keyword = split_vision_name(vision_name)
        unit = "PCS"
        if "unit" in columns:
            unit = str(ws.cell(excel_row, columns["unit"]).value or "PCS").strip() or "PCS"
        total_quantity = qty * purchase_quantity
        rows.append(
            BomRow(
                source_row=excel_row,
                index=len(rows) + 1,
                vision_name=vision_name,
                qty=qty,
                purchase_quantity=purchase_quantity,
                total_quantity=total_quantity,
                search_keyword=keyword,
                material_name_hint=material_hint,
                unit=unit,
            )
        )
    return rows


def prepare(args: argparse.Namespace) -> None:
    rows = read_bom_rows(args.bom, args.sheet, args.purchase_quantity)
    payload = {
        "source": str(args.bom),
        "sites": SITE_DEFAULTS,
        "items": [
            {
                "source_row": row.source_row,
                "index": row.index,
                "vision_name": row.vision_name,
                "search_keyword": row.search_keyword,
                "lcsc_keyword": row.vision_name,
                "lcsc_query_url": "https://so.szlcsc.com/global.html?k=" + quote_plus(row.vision_name),
                "ickey_search_url": (
                    "https://search.ickey.cn/?keyword="
                    + quote_plus(row.search_keyword)
                    + "&bom_ab=null"
                ),
                "qty": row.qty,
                "purchase_quantity": row.purchase_quantity,
                "total_quantity": row.total_quantity,
            }
            for row in rows
        ],
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(rows)} BOM query rows to {args.out}")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def lcsc_payload(keyword: str, current_page: int, page_size: int) -> dict[str, Any]:
    return {
        "currentPage": current_page,
        "pageSize": page_size,
        "catalogIdFilter": "",
        "brandIdFilter": "",
        "standardFilter": "",
        "arrangeFilter": "",
        "labelFilter": "",
        "authenticationFilter": "",
        "keyword": keyword,
        "sortNumber": 0,
        "satisfyStockType": "",
        "startPrice": "",
        "endPrice": "",
        "demandNumber": "",
        "spotFilter": 1,
        "discountFilter": 1,
        "hasDataFile": False,
        "brandPlaceFilter": "",
        "secondKeyword": "",
        "queryParameterValue": "",
        "lastParamName": "",
    }


def lcsc_headers(json_request: bool = False) -> dict[str, str]:
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125 Safari/537.36"
        ),
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
        "Referer": LCSC_HOME_URL,
    }
    if json_request:
        headers.update(
            {
                "Accept": "application/json, text/plain, */*",
                "Content-Type": "application/json;charset=UTF-8",
                "Origin": "https://so.szlcsc.com",
                "X-Requested-With": "XMLHttpRequest",
            }
        )
    else:
        headers["Accept"] = "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"
    return headers


def lcsc_opener(timeout: float) -> urllib.request.OpenerDirector:
    cookie_jar = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cookie_jar))
    request = urllib.request.Request(LCSC_HOME_URL, headers=lcsc_headers(json_request=False))
    with opener.open(request, timeout=timeout) as response:
        response.read(512)
    return opener


def lcsc_fetch_page(
    opener: urllib.request.OpenerDirector,
    keyword: str,
    current_page: int,
    page_size: int,
    timeout: float,
) -> dict[str, Any]:
    body = json.dumps(lcsc_payload(keyword, current_page, page_size), ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        LCSC_API_URL,
        data=body,
        headers=lcsc_headers(json_request=True),
        method="POST",
    )
    with opener.open(request, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8-sig"))


def lcsc_product_records(payload: dict[str, Any]) -> list[dict[str, Any]]:
    return list(
        payload.get("result", {})
        .get("searchResult", {})
        .get("productRecordList", [])
        or []
    )


def price_for_quantity(price_list: list[dict[str, Any]], total_quantity: float) -> float | None:
    fallback: float | None = None
    for price_record in price_list or []:
        price = to_number(price_record.get("productPrice") or price_record.get("thePrice"))
        if price is None:
            continue
        fallback = price if fallback is None else min(fallback, price)
        start = to_number(price_record.get("startPurchasedNumber") or price_record.get("spNumber")) or 0
        end = to_number(price_record.get("endPurchasedNumber") or price_record.get("epNumber"))
        if total_quantity >= start and (end is None or end < 0 or total_quantity <= end):
            return price
    return fallback


def first_number(*values: Any) -> float | None:
    for value in values:
        number = to_number(value)
        if number is not None:
            return number
    return None


def lcsc_stock_quantity(record: dict[str, Any], product: dict[str, Any]) -> float | None:
    return first_number(
        product.get("stockNumber"),
        record.get("totalStockNumber"),
        product.get("validStockNumber"),
        record.get("gdWarehouseStockNumber"),
        record.get("jsWarehouseStockNumber"),
        product.get("hkStockNumber"),
    )


def lcsc_in_stock(product: dict[str, Any], stock_qty: float | None) -> bool:
    if str(product.get("hasStockNow", "")).lower() == "yes":
        return True
    if str(product.get("productStockStatus", "")).lower() == "yes":
        return True
    return bool(stock_qty is not None and stock_qty > 0)


def lcsc_lead_time(product: dict[str, Any], in_stock: bool) -> str:
    if in_stock:
        return "\u73b0\u8d27"
    days = to_number(product.get("deliveryTimeWayDays"))
    if days is not None:
        return f"{int(days)}\u5929"
    weeks = to_number(product.get("agentProductReceiveWeeks"))
    if weeks is not None:
        return f"{int(weeks)}\u5468"
    receive_text = str(product.get("agentProductReceiveNumberText") or "").strip()
    return receive_text


def lcsc_offer_url(product: dict[str, Any], keyword: str) -> str:
    product_code = str(product.get("productCode") or "").strip()
    query = product_code or keyword
    return "https://so.szlcsc.com/global.html?k=" + quote_plus(query)


def lcsc_query_url(keyword: str) -> str:
    return "https://so.szlcsc.com/global.html?k=" + quote_plus(keyword)


def lcsc_matched_url(product: dict[str, Any], keyword: str) -> str:
    product_id = str(product.get("productId") or "").strip()
    if product_id:
        return "https://item.szlcsc.com/" + quote_plus(product_id) + ".html"
    return lcsc_offer_url(product, keyword)


def compact_join(parts: Iterable[Any], sep: str = ", ") -> str:
    return sep.join(str(part).strip() for part in parts if str(part or "").strip())


def lcsc_record_to_offer(
    record: dict[str, Any],
    total_quantity: float,
    keyword: str,
) -> dict[str, Any] | None:
    product = record.get("productVO") or {}
    if not product:
        return None
    stock_qty = lcsc_stock_quantity(record, product)
    in_stock = lcsc_in_stock(product, stock_qty)
    product_code = str(product.get("productCode") or record.get("lightProductCode") or "").strip()
    min_package = to_number(product.get("productMinEncapsulationNumber"))
    remark_parts = []
    if product_code:
        remark_parts.append(f"LCSC {product_code}")
    if min_package is not None:
        remark_parts.append(f"\u5305\u88c5\u6570\u91cf {int(min_package)}")
    unit_price = price_for_quantity(product.get("productPriceList") or [], total_quantity)
    if unit_price is None:
        remark_parts.append("\u672a\u89e3\u6790\u5230\u9636\u68af\u4ef7")
    return {
        "site": SITE_LCSC,
        "url": lcsc_query_url(keyword),
        "matched_url": lcsc_matched_url(product, keyword),
        "unit_price": unit_price,
        "stock_qty": stock_qty,
        "moq": first_number(product.get("minBuyNumber"), product.get("productMinEncapsulationNumber")),
        "lead_time": lcsc_lead_time(product, in_stock),
        "manufacturer": product.get("productGradePlateName") or record.get("lightBrandName") or "",
        "manufacturer_part_number": product.get("productModel") or record.get("lightProductModel") or "",
        "material_name": product.get("productType") or product.get("productName") or record.get("lightCatalogName") or "",
        "spec": compact_join(
            [
                product.get("productName"),
                product.get("encapsulationModel"),
                product.get("productModel"),
            ]
        ),
        "unit": product.get("productUnit") or "PCS",
        "in_stock": in_stock,
        "remark": "; ".join(remark_parts),
    }


def query_lcsc(args: argparse.Namespace) -> None:
    queries = read_json(args.queries)
    opener = lcsc_opener(args.timeout)
    output_items = []
    source_items = queries.get("items", [])
    if args.limit:
        source_items = source_items[: args.limit]

    for item in source_items:
        total_quantity = to_number(item.get("total_quantity")) or 0
        keywords = [str(item.get("lcsc_keyword") or item.get("vision_name") or "").strip()]
        spec_keyword = str(item.get("search_keyword") or "").strip()
        if args.fallback_spec and spec_keyword and spec_keyword not in keywords:
            keywords.append(spec_keyword)

        offers: list[dict[str, Any]] = []
        query_error = ""
        for keyword in [value for value in keywords if value]:
            for page in range(1, args.max_pages + 1):
                try:
                    payload = lcsc_fetch_page(opener, keyword, page, args.page_size, args.timeout)
                    records = lcsc_product_records(payload)
                    page_offers = [
                        offer
                        for offer in (lcsc_record_to_offer(record, total_quantity, keyword) for record in records)
                        if offer
                    ]
                    if page_offers:
                        offers.extend(page_offers)
                        break
                except Exception as exc:  # noqa: BLE001 - keep CLI useful for live site issues.
                    query_error = f"{type(exc).__name__}: {exc}"
                    break
            if offers:
                break
        output_item = {
            "source_row": item.get("source_row"),
            "vision_name": item.get("vision_name"),
            "offers": offers,
        }
        if query_error:
            output_item["query_error"] = query_error
        output_items.append(output_item)
        if args.delay:
            time.sleep(args.delay)

    result = {"sites": [SITE_LCSC], "items": output_items}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote LCSC offers for {len(output_items)} BOM rows to {args.out}")


def offer_in_stock(offer: dict[str, Any]) -> bool:
    explicit = normalize_bool(offer.get("in_stock"))
    if explicit is not None:
        return explicit
    stock = to_number(offer.get("stock_qty"))
    if stock is not None and stock > 0:
        return True
    text = " ".join(str(offer.get(k, "")) for k in ("lead_time", "remark", "stock_status")).lower()
    return any(token in text for token in IN_STOCK_TOKENS)


def offer_rank(offer: dict[str, Any], total_quantity: float) -> tuple[int, int, int, float]:
    stock = to_number(offer.get("stock_qty"))
    moq = to_number(offer.get("moq"))
    price = to_number(offer.get("unit_price"))
    return (
        0 if offer_in_stock(offer) else 1,
        0 if stock is not None and stock >= total_quantity else 1,
        0 if moq is not None and total_quantity >= moq else 1,
        price if price is not None else math.inf,
    )


def best_offer(offers: Iterable[dict[str, Any]], total_quantity: float) -> dict[str, Any] | None:
    candidates = [offer for offer in offers if offer]
    if not candidates:
        return None
    return sorted(candidates, key=lambda item: offer_rank(item, total_quantity))[0]


def index_results(results: dict[str, Any]) -> tuple[list[str], dict[int, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    sites = [str(site) for site in results.get("sites", []) if str(site).strip()]
    by_row: dict[int, list[dict[str, Any]]] = {}
    by_vision: dict[str, list[dict[str, Any]]] = {}
    for item in results.get("items", []):
        offers = list(item.get("offers", []) or [])
        source_row = item.get("source_row")
        if source_row is not None:
            try:
                by_row[int(source_row)] = offers
            except (TypeError, ValueError):
                pass
        vision_name = str(item.get("vision_name", "")).strip()
        if vision_name:
            by_vision[vision_name] = offers
        for offer in offers:
            site = str(offer.get("site", "")).strip()
            if site and site not in sites:
                sites.append(site)
    while len(sites) < 4:
        sites.append(f"Website{len(sites) + 1}")
    return sites[:4], by_row, by_vision


def safe_value(value: Any) -> Any:
    return "" if value is None else value


def price_value(offer: dict[str, Any] | None) -> Any:
    if not offer:
        return ""
    price = to_number(offer.get("unit_price"))
    return price if price is not None else safe_value(offer.get("unit_price"))


def fallback_quote_path(bom_path: Path) -> Path:
    return bom_path.with_name(bom_path.stem + "_" + DEFAULT_QUOTE_SHEET_NAME + bom_path.suffix)


def create_output_workbook(args: argparse.Namespace) -> tuple[Any, Any, Path]:
    if args.append_sheet:
        wb = load_workbook(args.bom)
        sheet_name = args.quote_sheet_name or DEFAULT_QUOTE_SHEET_NAME
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(title=sheet_name, index=1)
        return wb, ws, args.bom
    if not args.out:
        raise SystemExit("Pass --out, or use --append-sheet to write the quote sheet into the BOM workbook.")
    wb = Workbook()
    ws = wb.active
    ws.title = args.quote_sheet_name or "\u62a5\u4ef7\u5355"
    return wb, ws, args.out


def save_output_workbook(wb: Any, primary_path: Path, fallback_path: Path | None) -> Path:
    try:
        wb.save(primary_path)
        return primary_path
    except PermissionError:
        if fallback_path is None or fallback_path == primary_path:
            raise
        wb.save(fallback_path)
        return fallback_path
    except OSError:
        if fallback_path is None or fallback_path == primary_path:
            raise
        wb.save(fallback_path)
        return fallback_path


def build(args: argparse.Namespace) -> None:
    rows = read_bom_rows(args.bom, args.sheet, args.purchase_quantity)
    results = {}
    if args.results and args.results.exists():
        results = read_json(args.results)
    sites, offers_by_row, offers_by_vision = index_results(results)

    wb, ws, primary_output = create_output_workbook(args)

    quote_label = args.quote_label
    if not quote_label and args.purchase_quantity:
        quantity_text = int(args.purchase_quantity) if args.purchase_quantity.is_integer() else args.purchase_quantity
        quote_label = f"{quantity_text}\u5957"
    quote_label = quote_label or HEADER_PURCHASE_QTY_CN

    main_headers = [
        *MAIN_HEADERS_BASE,
        f"{quote_label}\u5355\u4ef7",
        HEADER_ORDER_QTY,
        f"{quote_label}\u603b\u4ef7",
        HEADER_REMARK,
        HEADER_LEAD_TIME,
    ]

    for col, header in enumerate(main_headers, start=1):
        cell = ws.cell(1, col, header)
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        cell.fill = MAIN_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    start_col = len(main_headers) + 1
    for site in sites:
        end_col = start_col + len(SITE_SUBHEADERS) - 1
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        site_cell = ws.cell(1, start_col, site)
        site_cell.fill = SITE_FILL
        site_cell.font = Font(bold=True)
        site_cell.alignment = Alignment(horizontal="center", vertical="center")
        for offset, subheader in enumerate(SITE_SUBHEADERS):
            cell = ws.cell(2, start_col + offset, subheader)
            cell.fill = SUB_FILL
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        start_col = end_col + 1

    for output_row, bom_row in enumerate(rows, start=3):
        offers = offers_by_row.get(bom_row.source_row) or offers_by_vision.get(bom_row.vision_name, [])
        site_best: dict[str, dict[str, Any] | None] = {}
        for site in sites:
            site_offers = [offer for offer in offers if str(offer.get("site", "")).strip() == site]
            site_best[site] = best_offer(site_offers, bom_row.total_quantity)
        selected = best_offer([offer for offer in site_best.values() if offer], bom_row.total_quantity)
        selected_price = to_number(selected.get("unit_price")) if selected else None
        selected_remark = selected.get("remark", "") if selected else WAIT_QUERY
        if selected:
            rank = offer_rank(selected, bom_row.total_quantity)
            warnings = []
            if rank[0]:
                warnings.append(WARN_NOT_STOCK)
            if rank[1]:
                warnings.append(WARN_STOCK_SHORT)
            if rank[2]:
                warnings.append(WARN_MOQ)
            if warnings:
                selected_remark = "; ".join([str(selected_remark)] + warnings).strip("; ")

        values = [
            bom_row.index,
            safe_value(selected.get("material_name") if selected else bom_row.material_name_hint),
            safe_value(selected.get("manufacturer") if selected else ""),
            safe_value(selected.get("manufacturer_part_number") if selected else ""),
            safe_value(selected.get("spec") if selected else bom_row.search_keyword),
            "",
            bom_row.qty,
            safe_value(selected.get("unit") if selected else bom_row.unit),
            selected_price if selected_price is not None else "",
            bom_row.total_quantity,
            None,
            selected_remark,
            safe_value(selected.get("lead_time") if selected else ""),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(output_row, col, value)
        price_col = get_column_letter(9)
        order_col = get_column_letter(10)
        ws.cell(output_row, 11, f"={price_col}{output_row}*{order_col}{output_row}")

        col = len(main_headers) + 1
        for site in sites:
            offer = site_best.get(site)
            site_values = [
                safe_value(offer.get("url") if offer else ""),
                safe_value((offer.get("matched_url") or offer.get("url")) if offer else ""),
                price_value(offer),
                safe_value(offer.get("stock_qty") if offer else ""),
                safe_value(offer.get("moq") if offer else ""),
                safe_value(offer.get("lead_time") if offer else ""),
            ]
            for value_index, value in enumerate(site_values):
                cell = ws.cell(output_row, col, value)
                if value_index in (0, 1) and isinstance(value, str) and value.startswith("http"):
                    cell.hyperlink = value
                    cell.style = "Hyperlink"
                col += 1

    apply_formatting(ws, max_row=ws.max_row, max_col=ws.max_column)
    fallback_path = args.out if args.out else fallback_quote_path(args.bom)
    primary_output.parent.mkdir(parents=True, exist_ok=True)
    if fallback_path:
        fallback_path.parent.mkdir(parents=True, exist_ok=True)
    saved_path = save_output_workbook(wb, primary_output, fallback_path)
    print(f"Wrote quote workbook to {saved_path}")


def apply_formatting(ws: Any, max_row: int, max_col: int) -> None:
    widths = {1: 8, 2: 18, 3: 18, 4: 24, 5: 36, 6: 18, 7: 10, 8: 10, 9: 14, 10: 18, 11: 16, 12: 34, 13: 16}
    for col in range(1, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 18)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(max_col)}{max_row}"
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = copy(THIN_BORDER)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row <= 2:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index in range(3, max_row + 1):
        ws.row_dimensions[row_index].height = 36
    for col in (7, 9, 10, 11):
        for row_index in range(3, max_row + 1):
            ws.cell(row_index, col).number_format = "#,##0.0000" if col == 9 else "#,##0"
    for col in range(14, max_col + 1):
        header = ws.cell(2, col).value
        if header == SITE_SUBHEADERS[2]:
            for row_index in range(3, max_row + 1):
                ws.cell(row_index, col).number_format = "#,##0.0000"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Prepare BOM queries and build quote workbooks.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    prepare_parser = subparsers.add_parser("prepare", help="Create query JSON from a BOM workbook.")
    prepare_parser.add_argument("--bom", type=Path, required=True, help="Input BOM .xlsx path.")
    prepare_parser.add_argument("--out", type=Path, required=True, help="Output query JSON path.")
    prepare_parser.add_argument("--sheet", help="Source sheet name. Defaults to raw BOM or first matching sheet.")
    prepare_parser.add_argument("--purchase-quantity", type=float, help="Default Purchase_Quantity when missing.")
    prepare_parser.set_defaults(func=prepare)

    lcsc_parser = subparsers.add_parser("query-lcsc", help="Query LCSC API from a prepared query JSON.")
    lcsc_parser.add_argument("--queries", type=Path, required=True, help="Query JSON from the prepare command.")
    lcsc_parser.add_argument("--out", type=Path, required=True, help="Output supplier offers JSON path.")
    lcsc_parser.add_argument("--page-size", type=int, default=30, help="LCSC page size. Default: 30.")
    lcsc_parser.add_argument("--max-pages", type=int, default=5, help="Only paginate when a page has no offers.")
    lcsc_parser.add_argument("--timeout", type=float, default=20.0, help="HTTP timeout in seconds.")
    lcsc_parser.add_argument("--delay", type=float, default=0.2, help="Delay between BOM row queries in seconds.")
    lcsc_parser.add_argument("--limit", type=int, help="Optional max BOM rows to query for testing.")
    lcsc_parser.add_argument("--fallback-spec", action="store_true", help="Try search_keyword when full Vision_name has no offers.")
    lcsc_parser.set_defaults(func=query_lcsc)

    build_parser = subparsers.add_parser("build", help="Build a quote workbook from BOM and offers JSON.")
    build_parser.add_argument("--bom", type=Path, required=True, help="Input BOM .xlsx path.")
    build_parser.add_argument("--out", type=Path, help="Output quote .xlsx path, or fallback path when using --append-sheet.")
    build_parser.add_argument("--results", type=Path, help="Supplier offers JSON path.")
    build_parser.add_argument("--sheet", help="Source sheet name. Defaults to raw BOM or first matching sheet.")
    build_parser.add_argument("--purchase-quantity", type=float, help="Default Purchase_Quantity when missing.")
    build_parser.add_argument("--quote-label", help="Header label such as 3000 sets or 6000 sets.")
    build_parser.add_argument("--append-sheet", action="store_true", help="Add/replace the quote sheet inside the BOM workbook.")
    build_parser.add_argument("--quote-sheet-name", default=DEFAULT_QUOTE_SHEET_NAME, help="Quote sheet name for new or appended output.")
    build_parser.set_defaults(func=build)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
